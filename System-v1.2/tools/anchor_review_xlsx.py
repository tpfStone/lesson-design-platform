"""把增量归纳结果转成 Excel 复核表（便于批注）

用法：
  python tools/anchor_review_xlsx.py            # 数学 + 英语各生成一份
  python tools/anchor_review_xlsx.py math       # 只生成数学
"""
import json
import sys
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 项目根目录
if BASE not in sys.path:
    sys.path.insert(0, BASE)

OUT_DIR = f"{BASE}/output/anchors"
BAND_CN = {"Q1": "低档(0-25%)", "Q2": "中低档(25-50%)", "Q3": "中高档(50-75%)", "Q4": "高档(75-100%)"}
HEADERS = ["维度", "档位", "状态", "新样本数", "归纳·优点", "归纳·缺点", "归纳·原文特征",
           "归纳·关键词", "原有·优点", "原有·缺点", "复核意见（请填写）"]
WIDTHS = [10, 16, 7, 8, 46, 60, 50, 40, 30, 34, 24]


def build(suffix: str) -> str:
    src = f"{OUT_DIR}/incremental_review_{suffix}.json"
    dst = f"{OUT_DIR}/incremental_review_{suffix}.xlsx"
    if not os.path.exists(src):
        print(f"⚠ 跳过 {suffix}：未找到 {src}")
        return ""
    d = json.load(open(src, encoding="utf-8"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{d['学科']}增量归纳"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F81BD")
        cell.alignment = Alignment(vertical="center", horizontal="center")

    rows = 0
    for dim, item in d["维度"].items():
        for band, s in item["归纳"].items():
            o = item.get("原有", {}).get(band) or {}
            ws.append([
                dim, BAND_CN.get(band, band), item["状态"].get(band, ""),
                item.get("样本数", {}).get(band, 0),
                s.get("优点", ""), s.get("缺点", ""), s.get("原文特征", ""),
                "、".join(s.get("关键词", []) or []),
                o.get("优点", ""), o.get("缺点", ""), "",
            ])
            rows += 1
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, rows + 2):
        for c in range(1, len(HEADERS) + 1):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(dst)
    print(f"✓ {dst}（{rows} 行）")
    return dst


if __name__ == "__main__":
    for sfx in (sys.argv[1:] or ["math", "english"]):
        build(sfx)
