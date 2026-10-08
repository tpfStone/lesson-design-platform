"""
解析"粗评归纳+模板示例.xlsx" → 结构化锚点库
输出: output/anchors/anchor_math.json / anchor_english.json
"""
import json, os, re
import openpyxl
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 项目根目录
if BASE not in sys.path:
    sys.path.insert(0, BASE)

SRC = r"E:/wechat/xwechat_files/wxid_gdc7waj7cony22_f188/msg/file/2026-09/粗评归纳+模板示例(1).xlsx"
OUT_DIR = f"{BASE}/output/anchors"
os.makedirs(OUT_DIR, exist_ok=True)

DIMS = ["教学内容", "教学理念", "学情分析", "教学目标", "教学资源", "教学方法",
        "过程设计", "课程思政", "课堂互动", "评价反馈", "版式观感", "文档规范"]
BAND_KEYS = ["Q1", "Q2", "Q3", "Q4"]


def parse_feature_text(txt: str) -> dict:
    """把单元格文本拆成 优点/缺点/原文特征"""
    t = str(txt).strip()
    if t in ("／", "/", ""):
        return {}
    res = {"raw": t}
    m = re.search(r"共性特征：?\s*(.*?)(?=原文特征|$)", t, re.S)
    feats = m.group(1) if m else t
    mo = re.search(r"优点：?\s*(.*?)(?=缺点|$)", feats, re.S)
    if mo:
        res["优点"] = mo.group(1).strip()
    md = re.search(r"缺点：?\s*(.*?)(?=$)", feats, re.S)
    if md:
        res["缺点"] = md.group(1).strip()
    if "优点" not in res and "缺点" not in res and feats.strip():
        res["特征"] = feats.strip()
    mr = re.search(r"原文特征[：:]?\s*(.*?)$", t, re.S)
    if mr:
        res["原文特征"] = mr.group(1).strip()
    return res


def parse_sheet(ws, subject: str) -> dict:
    # --- 主表：区间行 5/7/9/11（共性特征），列 2-14 = 总分+12维度 ---
    band_rows = {5: "Q1", 7: "Q2", 9: "Q3", 11: "Q4"}
    anchor = {"学科": subject, "标准": "competition", "维度": {}, "总分画像": {}}

    # 校验表头
    header = [ws.cell(4, c).value for c in range(2, 15)]
    assert "总分" in str(header[0]), f"{subject} 表头异常: {header[:3]}"

    for r, band in band_rows.items():
        # 总分（列2）
        txt = ws.cell(r, 2).value
        if txt and str(txt).strip() not in ("／", "/"):
            anchor["总分画像"][band] = parse_feature_text(txt)
        # 12 维度（列3..14）
        for i, dim in enumerate(DIMS):
            cell = ws.cell(r, 3 + i).value
            info = parse_feature_text(cell) if cell else {}
            if info:
                anchor["维度"].setdefault(dim, {})[band] = info

    # --- 模板库：从标题行后开始，行结构 [维度, 分数段, 示例模板, 返回主表] ---
    template_start = None
    for r in range(13, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v and "模板库" in str(v):
            template_start = r + 2  # 跳过标题行 + 表头行
            break
    if template_start:
        band_map = {"0–0.25": "Q1", "0.25–0.50": "Q2", "0.50–0.75": "Q3", "0.75–1.00": "Q4",
                    "0-0.25": "Q1", "0.25-0.50": "Q2", "0.50-0.75": "Q3", "0.75-1.00": "Q4"}
        for r in range(template_start, ws.max_row + 1):
            dim = ws.cell(r, 1).value
            band_raw = ws.cell(r, 2).value
            example = ws.cell(r, 3).value
            if not dim or not band_raw:
                continue
            band = band_map.get(str(band_raw).strip())
            if not band:
                continue
            dim = str(dim).strip()
            if dim == "总分":
                if example and str(example).strip() not in ("／", "/"):
                    anchor["总分画像"].setdefault(band, {})["模板示例"] = str(example).strip()
                continue
            if dim not in DIMS:
                continue
            if not example or str(example).strip() in ("／", "/"):
                continue
            anchor["维度"].setdefault(dim, {}).setdefault(band, {})["模板示例"] = str(example).strip()

    return anchor


wb = openpyxl.load_workbook(SRC, data_only=True)
for sn, fname in [("数学", "anchor_math.json"), ("英语", "anchor_english.json")]:
    a = parse_sheet(wb[sn], sn)
    # 统计
    filled = sum(len(v) for v in a["维度"].values())
    with_tpl = sum(1 for dim, bands in a["维度"].items() for b, d in bands.items() if d.get("模板示例"))
    path = os.path.join(OUT_DIR, fname)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(a, f, ensure_ascii=False, indent=1)
    print(f"✓ {sn}: 维度档位填充 {filled} 个（含模板 {with_tpl} 个）→ {path}")
    # 维度覆盖概览
    for dim in DIMS:
        bands = a["维度"].get(dim, {})
        print(f"    {dim:<6}: {', '.join(sorted(bands.keys())) if bands else '（无）'}")
